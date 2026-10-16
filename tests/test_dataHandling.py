import csv
import json

from openpyxl import load_workbook
import pytest

# @pytest.mark.dh
def test_jsonHandling():
    with open('testData\\creds.json') as data:
        formattedData = json.load(data)
        print(formattedData["positive"]["username"])


# @pytest.mark.dh
def test_csvHandling():
    finalData = []
    with open('testData\\credentails.csv') as data:
        formattedData = csv.DictReader(data)
        for i in formattedData:
            finalData.append(i)

    print(finalData[1]["username"])


def test_csvHandling():
    finalData = []
    with open('testData\\credentails.csv', mode='w',newline="") as data:
        formattedData = csv.DictWriter(data, fieldnames=["username","password"])
        formattedData.writeheader()
        formattedData.writerow({'username': 'tripur123_new', 'password': '1234_new'})

# @pytest.mark.dh
def test_excel():
    values=[]
    workBook = load_workbook("testData\\sample_creds.xlsx")
    sheet = workBook['Sheet2']
    for i in sheet.iter_rows(min_row=2,values_only=True):
        values.append(i)

    print(values)

@pytest.mark.dh
def test_excel():
    # values=[]
    workBook = load_workbook("testData\\sample_creds.xlsx")
    sheet = workBook['Sheet2']
    # sheet["A2"]="hello"
    # sheet.append(["test_new","newData"])
    sheet.delete_rows(3,3)
    workBook.save("testData\\sample_creds.xlsx")



